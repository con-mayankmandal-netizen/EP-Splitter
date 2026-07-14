import os
import tempfile
import unittest
import sys
from hashlib import sha256
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
from main import analyse_docx, build_docx


def body_signatures(path: str) -> list[str]:
    """Hash every non-section direct body element in document order."""
    doc = Document(path)
    return [
        sha256(etree.tostring(child, with_tail=False)).hexdigest()
        for child in doc.element.body.iterchildren()
        if child.tag != qn("w:sectPr")
    ]


def build_fixture(path: str) -> None:
    doc = Document()
    doc.add_heading("Fixture novel", level=1)
    for number in range(1, 4):
        heading = f"Chapter {number} Unpunctuated title" if number == 2 else f"Chapter {number}: Chapter {number}"
        doc.add_heading(heading, level=1)
        doc.add_paragraph(f"Chapter {number} Repeated running title")
        doc.add_paragraph(f"Unique chapter {number} content.")
    doc.add_heading("Chapter Summary", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Number"
    table.cell(0, 1).text = "Chapter"
    for number in range(1, 4):
        cells = table.add_row().cells
        cells[0].text = str(number)
        cells[1].text = f"Chapter {number}"
    doc.add_heading("Missing / Empty Chapters", level=1)
    doc.add_paragraph("None")
    doc.save(path)


class SplitterRegressionTests(unittest.TestCase):
    def assert_lossless_individual_split(self, source_path: str, expected_count: int) -> None:
        analysis = analyse_docx(source_path)
        episodes = analysis["episodes"]
        self.assertEqual(expected_count, len(episodes))
        self.assertEqual(list(range(1, expected_count + 1)), [e["number"] for e in episodes])

        with tempfile.TemporaryDirectory() as temp_dir:
            combined_signatures: list[str] = []
            for index in range(len(episodes)):
                output_path = str(Path(temp_dir) / f"episode-{index + 1}.docx")
                build_docx(source_path, episodes, [index], output_path)
                combined_signatures.extend(body_signatures(output_path))

            self.assertEqual(body_signatures(source_path), combined_signatures)

    def test_repeated_titles_and_summary_table_are_not_boundaries(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = str(Path(temp_dir) / "fixture.docx")
            build_fixture(source_path)
            self.assert_lossless_individual_split(source_path, 3)

    @unittest.skipUnless(
        os.environ.get("EP_SPLITTER_TEST_DOCX"),
        "Set EP_SPLITTER_TEST_DOCX to run against a real source document.",
    )
    def test_supplied_document_is_100_chapters_and_lossless(self):
        self.assert_lossless_individual_split(os.environ["EP_SPLITTER_TEST_DOCX"], 100)


if __name__ == "__main__":
    unittest.main()
